#!/usr/bin/env python3
"""
Messy Data Cleaning Pipeline: Dedupes, Standardizes, Validates
==============================================================

Takes a messy customer CSV and produces:

    output/clean_customers.csv    the cleaned, deduplicated dataset
    output/change_log.csv         every single edit, traceable to its source row
    output/rejected_rows.csv      rows with no usable contact info, with the reason
    output/cleaning_report.xlsx   all of the above plus a summary, in one workbook

What it fixes
-------------
* Placeholders ("N/A", "null", "-") become real blanks
* Names: "SMITH, JOHN" -> "John Smith", stray digits/symbols removed, McDonald/O'Brien handled
* Emails: lowercased, trimmed, common domain typos fixed (gmial.com -> gmail.com), invalid ones removed
* Phones: any format -> (416) 555-0187; invalid numbers removed
* Province/state: "Ontario", "ont.", "on" -> ON
* Dates: five input formats -> YYYY-MM-DD (slashed dates are read as MM/DD/YYYY)
* Amounts: "$1,234.50" / "CAD 120" -> numbers; negatives and junk removed
* Duplicates: rows are linked by customer ID, email, or name+phone; the most complete
  row is kept and any blanks in it are filled from its duplicates

Usage
-----
    python generate_messy_data.py        # make sample data first
    python clean_data.py
    python clean_data.py --input my_file.csv --output-dir results
"""

import argparse
import logging
import re
import sys
import unicodedata
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

log = logging.getLogger("cleaner")

DATA_COLS = ["customer_id", "full_name", "email", "phone", "city",
             "province_state", "signup_date", "total_spend"]
LOG_FILE = Path("cleaning.log")

MISSING_TOKENS = {"", "n/a", "na", "null", "none", "nan", "-", "--", "unknown", "?"}


def is_missing(value) -> bool:
    if value is None:
        return True
    try:
        if pd.isna(value):
            return True
    except (TypeError, ValueError):
        pass
    return str(value).strip().lower() in MISSING_TOKENS


def squash(text) -> str:
    """Collapse runs of whitespace and trim the ends."""
    return re.sub(r"\s+", " ", str(text)).strip()


def strip_accents(text: str) -> str:
    return "".join(ch for ch in unicodedata.normalize("NFKD", text)
                   if not unicodedata.combining(ch))


# --------------------------------------------------------------------------
# Field cleaners: each takes a non-blank raw value and returns
# (new_value or None, short description of what was done)
# --------------------------------------------------------------------------
def smart_title(word: str) -> str:
    """Title-case a word while respecting O'Brien, McDonald and Jean-Luc."""
    def cap(seg: str) -> str:
        seg = seg.lower()
        if seg.startswith("mc") and len(seg) > 2:
            return "Mc" + seg[2:].capitalize()
        return seg.capitalize()

    return "-".join("'".join(cap(p) for p in hyphen_part.split("'"))
                    for hyphen_part in word.split("-"))


def clean_id(value):
    return squash(value).upper(), "Standardized customer ID"


def clean_name(value):
    text = squash(value)
    action = "Fixed capitalization / spacing"

    if "," in text:
        last, _, first = text.partition(",")
        if last.strip() and first.strip():
            text = f"{first.strip()} {last.strip()}"
            action = "Reordered 'Last, First'"

    stripped = squash(re.sub(r"[^A-Za-z\u00C0-\u017F\s'\-.]", "", text))
    if stripped != text:
        action = "Removed invalid characters from name"
    if not stripped:
        return None, "Removed invalid name"

    return " ".join(smart_title(w) for w in stripped.split(" ")), action


EMAIL_RE = re.compile(r"^[a-z0-9._%+\-]+@[a-z0-9\-]+(\.[a-z0-9\-]+)*\.[a-z]{2,}$")
DOMAIN_FIXES = {
    "gmial.com": "gmail.com", "gmai.com": "gmail.com", "gamil.com": "gmail.com",
    "gmail.con": "gmail.com", "gmail.cmo": "gmail.com", "gmail.co": "gmail.com",
    "hotmial.com": "hotmail.com", "hotmail.con": "hotmail.com",
    "yaho.com": "yahoo.com", "yahooo.com": "yahoo.com", "yahoo.con": "yahoo.com",
    "outlok.com": "outlook.com", "outlook.con": "outlook.com",
    "icloud.con": "icloud.com",
}


def clean_email(value):
    text = re.sub(r"\s+", "", str(value)).lower()
    local, at, domain = text.rpartition("@")
    action = "Standardized email (lowercase / trim)"

    if at and domain in DOMAIN_FIXES:
        text = f"{local}@{DOMAIN_FIXES[domain]}"
        action = "Fixed email domain typo"

    if not EMAIL_RE.match(text) or ".." in text:
        return None, "Removed invalid email"
    return text, action


def clean_phone(value):
    text = re.sub(r"(?i)\s*(?:ext\.?|x)\s*\d+\s*$", "", str(value))  # drop extensions
    digits = re.sub(r"\D", "", text)
    if len(digits) == 11 and digits.startswith("1"):
        digits = digits[1:]
    # North American numbers: 10 digits, area code and exchange can't start with 0 or 1
    if len(digits) != 10 or digits[0] in "01" or digits[3] in "01":
        return None, "Removed invalid phone number"
    return f"({digits[:3]}) {digits[3:6]}-{digits[6:]}", "Standardized phone format"


US_STATES = {
    "alabama": "AL", "alaska": "AK", "arizona": "AZ", "arkansas": "AR",
    "california": "CA", "colorado": "CO", "connecticut": "CT", "delaware": "DE",
    "district of columbia": "DC", "florida": "FL", "georgia": "GA", "hawaii": "HI",
    "idaho": "ID", "illinois": "IL", "indiana": "IN", "iowa": "IA", "kansas": "KS",
    "kentucky": "KY", "louisiana": "LA", "maine": "ME", "maryland": "MD",
    "massachusetts": "MA", "michigan": "MI", "minnesota": "MN", "mississippi": "MS",
    "missouri": "MO", "montana": "MT", "nebraska": "NE", "nevada": "NV",
    "new hampshire": "NH", "new jersey": "NJ", "new mexico": "NM", "new york": "NY",
    "north carolina": "NC", "north dakota": "ND", "ohio": "OH", "oklahoma": "OK",
    "oregon": "OR", "pennsylvania": "PA", "rhode island": "RI",
    "south carolina": "SC", "south dakota": "SD", "tennessee": "TN", "texas": "TX",
    "utah": "UT", "vermont": "VT", "virginia": "VA", "washington": "WA",
    "west virginia": "WV", "wisconsin": "WI", "wyoming": "WY",
}
CA_PROVINCES = {
    "alberta": "AB", "british columbia": "BC", "manitoba": "MB",
    "new brunswick": "NB", "newfoundland and labrador": "NL", "nova scotia": "NS",
    "northwest territories": "NT", "nunavut": "NU", "ontario": "ON",
    "prince edward island": "PE", "quebec": "QC", "saskatchewan": "SK", "yukon": "YT",
}
REGION_MAP = {name: code for name, code in {**US_STATES, **CA_PROVINCES}.items()}
REGION_MAP.update({code.lower(): code for code in set(REGION_MAP.values())})
REGION_MAP.update({
    "ont": "ON", "que": "QC", "alta": "AB", "man": "MB", "sask": "SK",
    "nfld": "NL", "pei": "PE", "newfoundland": "NL", "b c": "BC",
})


def clean_region(value):
    key = re.sub(r"[^a-z ]", "", strip_accents(str(value)).lower())
    key = squash(key)
    code = REGION_MAP.get(key)
    if code is None:
        return None, "Removed unrecognized province/state"
    return code, "Standardized province/state code"


def clean_city(value):
    text = squash(re.sub(r"[^A-Za-z\u00C0-\u017F\s'\-.]", "", str(value)))
    if not text:
        return None, "Removed invalid city"
    return " ".join(smart_title(w) for w in text.split(" ")), "Fixed capitalization / spacing"


DATE_FORMATS = ["%Y-%m-%d", "%m/%d/%Y", "%b %d, %Y", "%d-%b-%Y", "%B %d, %Y"]


def clean_date(value):
    text = squash(value)
    for fmt in DATE_FORMATS:
        try:
            parsed = datetime.strptime(text, fmt)
        except ValueError:
            continue
        if parsed > datetime.now():
            return None, "Removed future date"
        return parsed.strftime("%Y-%m-%d"), "Standardized date to YYYY-MM-DD"
    return None, "Removed unparseable date"


def clean_spend(value):
    text = re.sub(r"(?i)cad|usd|[$,\s]", "", str(value))
    try:
        amount = float(text)
    except ValueError:
        return None, "Removed non-numeric amount"
    if amount < 0:
        return None, "Removed negative amount"
    return round(amount, 2), "Standardized amount"


COLUMN_CLEANERS = {
    "customer_id": clean_id,
    "full_name": clean_name,
    "email": clean_email,
    "phone": clean_phone,
    "city": clean_city,
    "province_state": clean_region,
    "signup_date": clean_date,
    "total_spend": clean_spend,
}


# --------------------------------------------------------------------------
# Pipeline steps
# --------------------------------------------------------------------------
def load_raw(path: Path) -> pd.DataFrame:
    if not path.exists():
        sys.exit(f"Input file not found: {path}\n"
                 "Run 'python generate_messy_data.py' first to create sample data.")
    df = pd.read_csv(path, dtype=str, keep_default_na=False)
    missing = [c for c in DATA_COLS if c not in df.columns]
    if missing:
        sys.exit(f"Input file is missing required columns: {', '.join(missing)}")
    df = df[DATA_COLS].copy()
    df.insert(0, "source_row", range(2, len(df) + 2))  # line number in the CSV file
    return df


def _same(raw, new) -> bool:
    if new is None:
        return str(raw) == ""
    if isinstance(new, float):
        try:
            return float(str(raw).strip()) == new
        except ValueError:
            return False
    return str(raw) == str(new)


def standardize(df: pd.DataFrame, changes: list) -> pd.DataFrame:
    """Apply every field cleaner and log each edit."""
    out = df.copy()
    for col, cleaner in COLUMN_CLEANERS.items():
        new_values = []
        for source_row, raw in zip(out["source_row"], out[col]):
            if is_missing(raw):
                new = None
                action = "Replaced placeholder with blank" if str(raw).strip() else None
            else:
                new, action = cleaner(raw)
            if action and not _same(raw, new):
                changes.append({
                    "source_row": source_row, "column": col,
                    "original_value": raw,
                    "new_value": "" if new is None else new,
                    "action": action,
                })
            new_values.append(new)
        out[col] = new_values
    out["total_spend"] = pd.to_numeric(out["total_spend"], errors="coerce")
    return out


def group_ids(df: pd.DataFrame, conflicts: list = None) -> list:
    """
    Link rows that share a customer ID, an email, or a name + phone.

    Safety rule: rows that share an ID or email but have *different* names are
    NOT merged (could be family members sharing an email). They are reported in
    `conflicts` for human review instead.
    """
    parent = list(range(len(df)))
    names = [n.lower() if pd.notna(n) else None for n in df["full_name"]]

    def find(i):
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i

    def compatible(a, b):
        return names[a] is None or names[b] is None or names[a] == names[b]

    seen = {}
    for pos, (cid, email, name, phone) in enumerate(
            zip(df["customer_id"], df["email"], df["full_name"], df["phone"])):
        keys = []
        if pd.notna(cid):
            keys.append("id:" + cid)
        if pd.notna(email):
            keys.append("email:" + email)
        if pd.notna(name) and pd.notna(phone):
            keys.append("np:" + name.lower() + "|" + phone)
        for key in keys:
            if key not in seen:
                seen[key] = pos
                continue
            other = seen[key]
            if compatible(pos, other):
                ra, rb = find(pos), find(other)
                if ra != rb:
                    parent[rb] = ra
            elif conflicts is not None:
                conflicts.append((pos, other, key))
    return [find(i) for i in range(len(df))]


def deduplicate(df: pd.DataFrame, changes: list) -> pd.DataFrame:
    """Keep the most complete row per customer and fill its gaps from duplicates."""
    work = df.reset_index(drop=True)
    rows_by_pos = work["source_row"].tolist()
    conflicts = []
    work["_group"] = group_ids(work, conflicts)
    for pos, other, key in conflicts:
        kind, _, value = key.partition(":")
        changes.append({
            "source_row": int(rows_by_pos[pos]), "column": "(entire row)",
            "original_value": f"{kind}: {value}",
            "new_value": f"Not merged with row {int(rows_by_pos[other])}",
            "action": "Review: shared ID/email but different names",
        })
    work["_score"] = work[DATA_COLS].notna().sum(axis=1)
    work = work.sort_values(["_score", "signup_date", "source_row"],
                            ascending=[False, False, True])

    top = work.drop_duplicates("_group").set_index("_group")
    merged = work.groupby("_group", sort=False)[DATA_COLS + ["source_row"]].first()
    merged = merged.loc[top.index]

    for col in DATA_COLS:
        filled = top[col].isna() & merged[col].notna()
        for group in top.index[filled]:
            changes.append({
                "source_row": int(top.at[group, "source_row"]), "column": col,
                "original_value": "", "new_value": merged.at[group, col],
                "action": "Filled blank from duplicate row",
            })

    kept_row = top["source_row"].to_dict()
    for _, row in work[work.duplicated("_group", keep="first")].iterrows():
        label = " / ".join(str(v) for v in (row["full_name"], row["email"])
                           if pd.notna(v)) or "(no name or email)"
        changes.append({
            "source_row": int(row["source_row"]), "column": "(entire row)",
            "original_value": label,
            "new_value": f"Merged into row {int(kept_row[row['_group']])}",
            "action": "Removed duplicate",
        })

    merged["source_row"] = top["source_row"]
    return merged.reset_index(drop=True).sort_values("source_row").reset_index(drop=True)


def split_rejected(df: pd.DataFrame, raw: pd.DataFrame, changes: list):
    """Rows with neither a valid email nor a valid phone can't be used."""
    no_contact = df["email"].isna() & df["phone"].isna()
    rejected_rows = set(df.loc[no_contact, "source_row"])

    rejected = raw[raw["source_row"].isin(rejected_rows)].copy()
    rejected["reject_reason"] = "No valid email or phone number"

    for row in rejected_rows:
        changes.append({
            "source_row": int(row), "column": "(entire row)",
            "original_value": "", "new_value": "Moved to rejected_rows",
            "action": "Rejected: no valid email or phone",
        })
    return df[~no_contact].reset_index(drop=True), rejected


# --------------------------------------------------------------------------
# Reporting
# --------------------------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def style_sheet(ws, widths: dict, money_col=None):
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = width
    if money_col:
        for row in ws.iter_rows(min_row=2, min_col=money_col, max_col=money_col):
            row[0].number_format = '"$"#,##0.00'
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def build_summary(n_in, n_out, n_rejected, changes_df) -> pd.DataFrame:
    n_dups = int((changes_df["action"] == "Removed duplicate").sum()) if len(changes_df) else 0
    n_review = int(changes_df["action"].str.startswith("Review:").sum()) if len(changes_df) else 0
    cell_fixes = changes_df[~changes_df["column"].eq("(entire row)")] if len(changes_df) else changes_df
    metrics = [
        ("Rows in (messy file)", n_in),
        ("Rows out (clean file)", n_out),
        ("Duplicate rows removed", n_dups),
        ("Rows rejected (no usable contact)", n_rejected),
        ("Possible conflicts flagged for review", n_review),
        ("Individual cell fixes", len(cell_fixes)),
        ("", ""),
        ("Cell fixes by type", ""),
    ]
    if len(cell_fixes):
        for action, count in cell_fixes["action"].value_counts().items():
            metrics.append((f"   {action}", int(count)))
    return pd.DataFrame(metrics, columns=["Metric", "Value"])


def export_excel(path: Path, summary, clean, changes_df, rejected):
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False)
        clean.to_excel(writer, sheet_name="Clean Data", index=False)
        changes_df.to_excel(writer, sheet_name="Change Log", index=False)
        rejected.to_excel(writer, sheet_name="Rejected Rows", index=False)

        style_sheet(writer.sheets["Summary"], {1: 42, 2: 14})
        for row in writer.sheets["Summary"].iter_rows(min_row=2):
            row[1].alignment = Alignment(horizontal="left")
        style_sheet(writer.sheets["Clean Data"],
                    {1: 14, 2: 24, 3: 34, 4: 18, 5: 16, 6: 16, 7: 14, 8: 14, 9: 12},
                    money_col=DATA_COLS.index("total_spend") + 1)
        style_sheet(writer.sheets["Change Log"], {1: 12, 2: 16, 3: 34, 4: 34, 5: 40})
        style_sheet(writer.sheets["Rejected Rows"],
                    {1: 12, 2: 14, 3: 24, 4: 30, 5: 18, 6: 16, 7: 16, 8: 14, 9: 14, 10: 34})
        if rejected.empty:
            writer.sheets["Rejected Rows"]["A2"] = "No rejected rows."


# --------------------------------------------------------------------------
# Orchestration
# --------------------------------------------------------------------------
def run_pipeline(input_path: Path, output_dir: Path) -> dict:
    raw = load_raw(input_path)
    log.info("Loaded %d rows from %s", len(raw), input_path)

    changes = []
    standardized = standardize(raw, changes)
    log.info("Standardized fields: %d cell edits so far", len(changes))

    deduped = deduplicate(standardized, changes)
    log.info("Deduplicated: %d rows -> %d", len(standardized), len(deduped))

    clean, rejected = split_rejected(deduped, raw, changes)
    log.info("Rejected %d rows with no usable contact info", len(rejected))

    changes_df = pd.DataFrame(
        changes, columns=["source_row", "column", "original_value", "new_value", "action"]
    ).sort_values(["source_row", "column"]).reset_index(drop=True)

    clean_out = clean[DATA_COLS + ["source_row"]]
    summary = build_summary(len(raw), len(clean_out), len(rejected), changes_df)

    output_dir.mkdir(parents=True, exist_ok=True)
    clean_out.to_csv(output_dir / "clean_customers.csv", index=False, float_format="%.2f")
    changes_df.to_csv(output_dir / "change_log.csv", index=False)
    rejected.to_csv(output_dir / "rejected_rows.csv", index=False)
    export_excel(output_dir / "cleaning_report.xlsx", summary, clean_out, changes_df, rejected)

    n_dups = int((changes_df["action"] == "Removed duplicate").sum())
    return {
        "input_rows": len(raw), "output_rows": len(clean_out),
        "duplicates_removed": n_dups, "rejected": len(rejected),
        "changes": changes_df, "clean": clean_out, "summary": summary,
    }


def setup_logging() -> None:
    fmt = logging.Formatter("%(asctime)s  %(levelname)-7s  %(message)s", "%H:%M:%S")
    console = logging.StreamHandler(sys.stdout)
    console.setFormatter(fmt)
    file_handler = logging.FileHandler(LOG_FILE, encoding="utf-8")
    file_handler.setFormatter(fmt)
    log.setLevel(logging.INFO)
    log.addHandler(console)
    log.addHandler(file_handler)


def main() -> int:
    parser = argparse.ArgumentParser(description="Clean a messy customer CSV")
    parser.add_argument("--input", type=Path, default=Path("data/messy_customers.csv"))
    parser.add_argument("--output-dir", type=Path, default=Path("output"))
    args = parser.parse_args()

    setup_logging()
    result = run_pipeline(args.input, args.output_dir)

    cell_fixes = int((~result["changes"]["column"].eq("(entire row)")).sum())
    log.info("-" * 52)
    log.info("DONE: %d messy rows -> %d clean records", result["input_rows"], result["output_rows"])
    log.info("  Duplicates removed : %d", result["duplicates_removed"])
    log.info("  Rows rejected      : %d", result["rejected"])
    log.info("  Cell-level fixes   : %d", cell_fixes)
    log.info("Report saved to: %s", args.output_dir / "cleaning_report.xlsx")
    return 0


if __name__ == "__main__":
    sys.exit(main())
